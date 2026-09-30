# -*- coding: utf-8 -*-
"""확장이 찾아온 주소를 엑셀에 합치고, 수집용 주소 목록을 뽑는다.

하는 일
-------
1. 검색결과.json 을 읽어 '확실'만 엑셀의 쿠팡URL 칸에 채운다
2. '애매'는 채우지 않고 따로 뽑아 사람이 보게 한다
3. URL이 채워진 항목으로 '수집용_주소목록.txt' 를 만든다

'애매'를 자동으로 채우지 않는 이유
--------------------------------
이름이 같아도 용량이 다른 제품이 흔하다(2.1L 1개 vs 4.2L 2개).
잘못 채우면 엉뚱한 제품의 상세를 긁어 놓고도 맞다고 믿게 된다.
그 오염은 나중에 찾아내기가 거의 불가능하다. 그래서 확실한 것만 채운다.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

PID_RE = re.compile(r"/v[pm]/products/(\d+)")


def merge(xlsx_path, found_path, out_xlsx, out_urls, out_review):
    found = json.loads(Path(found_path).read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(xlsx_path)
    stats = {"확실": 0, "애매": 0, "못찾음": 0, "이미있음": 0, "기록없음": 0}
    review, urls = [], []

    for ws in wb.worksheets:
        if ws.title == "요약":
            continue
        hdr = [c.value for c in ws[1]]
        if "카테고리" not in hdr or "순위" not in hdr:
            continue
        if "쿠팡URL" not in hdr:
            ws.cell(row=1, column=len(hdr) + 1, value="쿠팡URL")
            hdr.append("쿠팡URL")
        cu = hdr.index("쿠팡URL") + 1
        cc, cr = hdr.index("카테고리") + 1, hdr.index("순위") + 1
        cn = hdr.index("제품명") + 1

        for row in range(2, ws.max_row + 1):
            cat, rank = ws.cell(row, cc).value, ws.cell(row, cr).value
            if not cat or rank is None:
                continue
            key = f"{cat}-{int(rank):02d}"
            cur = ws.cell(row, cu).value
            if cur and PID_RE.search(str(cur)):
                if ws.title == "전체":
                    stats["이미있음"] += 1
                    urls.append((key, str(cur), ws.cell(row, cn).value))
                continue
            rec = found.get(key)
            if not rec:
                if ws.title == "전체":
                    stats["기록없음"] += 1
                continue
            v = rec.get("verdict")
            if v == "확실" and rec.get("url"):
                ws.cell(row, cu, rec["url"])
                if ws.title == "전체":
                    stats["확실"] += 1
                    urls.append((key, rec["url"], ws.cell(row, cn).value))
            elif ws.title == "전체":
                stats[v if v in stats else "못찾음"] += 1
                review.append({
                    "항목": key,
                    "찾던 제품": rec.get("name", ""),
                    "기대가": rec.get("wantPrice"),
                    "판정": v,
                    "후보": [
                        {"딜번호": c.get("pid"), "제품명": c.get("name"),
                         "가격": c.get("price"),
                         "이름일치": c.get("nameScore"),
                         "가격차": (None if c.get("priceGap") is None
                                  else round(c["priceGap"] * 100)),
                         "광고": c.get("isAd"),
                         "주소": c.get("url")}
                        for c in (rec.get("candidates") or [])[:5]
                    ],
                })

    wb.save(out_xlsx)
    Path(out_urls).write_text(
        "\n".join(f"{k} | {u}" for k, u, _ in urls), encoding="utf-8")
    Path(out_review).write_text(
        json.dumps(review, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"엑셀 저장        : {out_xlsx}")
    print(f"수집용 주소목록  : {out_urls}  ({len(urls)}개)")
    print(f"사람이 볼 것     : {out_review}  ({len(review)}개)")
    print()
    for k, v in stats.items():
        print(f"   {k:<8} {v}")
    return stats


if __name__ == "__main__":
    merge(sys.argv[1], sys.argv[2],
          sys.argv[3] if len(sys.argv) > 3 else "쿠팡_제품목록_URL반영.xlsx",
          sys.argv[4] if len(sys.argv) > 4 else "수집용_주소목록.txt",
          sys.argv[5] if len(sys.argv) > 5 else "확인필요_후보.json")
