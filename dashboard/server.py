"""Read-only local dashboard: python dashboard/server.py --port 3000."""
from __future__ import annotations

import argparse
from collections import Counter
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
import json
import math
from pathlib import Path
import sys
import re
from urllib.parse import quote, urlsplit

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "thumbnail_analyzer"))
from analyze_thumbnails import ProductInfo, row_from_info, parse_capacity, unit_price_basis  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from pydantic import ValidationError  # noqa: E402

CATEGORIES = ["세탁세제", "캡슐세제", "섬유유연제", "섬유탈취제", "주방세제", "살균소독제"]
SURVEY_DATE = "2026-09-21"
MAX_BODY = 4 * 1024 * 1024


def normalize(category, entries, survey_date=SURVEY_DATE):
    if category not in CATEGORIES:
        raise ValueError("지원하지 않는 카테고리입니다.")
    if not isinstance(entries, list) or not 1 <= len(entries) <= 1000:
        raise ValueError("JSON은 1~1,000개 항목의 배열이어야 합니다.")
    products = []
    for index, entry in enumerate(entries):
        if not isinstance(entry, dict):
            raise ValueError(f"{index + 1}번째 항목은 JSON 객체여야 합니다.")
        entry = dict(entry)
        filename = entry.pop("file", None) or f"{index + 1:02}.jpg"
        if not isinstance(filename, str) or len(filename) > 500:
            raise ValueError(f"{index + 1}번째 파일명이 올바르지 않습니다.")
        try:
            info = ProductInfo.model_validate(entry, strict=True)
        except ValidationError as exc:
            error = exc.errors()[0]
            raise ValueError(f"{index + 1}번째 항목 · {'.'.join(map(str, error['loc']))}: {error['msg']}") from exc
        for key in ("price_krw", "review_count"):
            value = getattr(info, key)
            if value is not None and not 0 <= value <= 10**12:
                raise ValueError(f"{index + 1}번째 항목 · {key}: 0 이상의 유효한 숫자를 입력하세요.")
        for value in info.model_dump().values():
            if isinstance(value, str) and (len(value) > 10000 or re.search(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", value)):
                raise ValueError(f"{index + 1}번째 항목의 텍스트가 너무 길거나 허용되지 않는 제어 문자가 포함되어 있습니다.")
        row = row_from_info(info, filename, {"survey_date": survey_date, "channel": "쿠팡", "category": category})
        if any(isinstance(v, float) and not math.isfinite(v) for v in row.values()):
            raise ValueError(f"{index + 1}번째 항목의 계산 결과가 유효하지 않습니다.")
        products.append({
            **info.model_dump(), "id": f"{category}-{index + 1}", "file": filename,
            "category": category, "rank": row["순위"], "unit_price": row["단위당가격(원)"],
            "basis": row["단가기준"], "total_capacity": row["총용량"], "capacity_unit": row["단위"],
            "validation": row["단가검증"], "buyers": row["월구매자수(명)"],
            "revenue": row["추정월매출(원)"], "search_url": row["검색링크"], "row": row,
        })
    return products


def load_products():
    return [p for cat in CATEGORIES for p in normalize(cat, json.loads((ROOT / "data/extracted" / f"{cat}.json").read_text(encoding="utf-8")))]


def load_specs():
    return json.loads((ROOT / "문서생성/견적요청서_데이터.json").read_text(encoding="utf-8"))


def summarize(products):
    counts = Counter(p["validation"] for p in products)
    categories = []
    for cat in CATEGORIES:
        group = [p for p in products if p["category"] == cat]
        categories.append({"name": cat, "count": len(group), "revenue": sum(p["revenue"] or 0 for p in group),
                           "known_revenue": sum(p["revenue"] is not None for p in group),
                           "brands": len({p["brand"] for p in group if p["brand"]}),
                           "roles": dict(Counter(p["product_role"] for p in group))})
    return {"count": len(products), "brands": len({p["brand"] for p in products if p["brand"]}),
            "revenue": sum(p["revenue"] or 0 for p in products), "validation": dict(counts),
            "verified_percent": round(counts["일치"] / len(products) * 100, 1) if products else 0,
            "roles": dict(Counter(p["product_role"] for p in products)), "categories": categories}


def snapshot(products, imported=False):
    return {"products": products, "summary": summarize(products), "specs": load_specs(),
            "survey_date": SURVEY_DATE, "channel": "쿠팡", "imported": imported,
            "notice": "추정월매출 = 월구매자수 × 판매가 × 2. 조사 표본의 규모 비교용이며 실제 시장 전체 매출이 아닙니다."}


def products_from_payload(payload):
    if not isinstance(payload, dict):
        raise ValueError("JSON 객체가 필요합니다.")
    overrides = payload.get("overrides", {})
    if not isinstance(overrides, dict) or any(c not in CATEGORIES for c in overrides):
        raise ValueError("카테고리 데이터를 확인하세요.")
    products = []
    for cat in CATEGORIES:
        if cat in overrides:
            products.extend(normalize(cat, overrides[cat], "업로드 데이터 (조사일 미확인)"))
        else:
            products.extend(normalize(cat, json.loads((ROOT / "data/extracted" / f"{cat}.json").read_text(encoding="utf-8"))))
    return products


def make_workbook(products):
    book = Workbook()
    sheet = book.active
    sheet.title = "전체"
    headers = list(products[0]["row"])
    sheet.append(headers)
    for product in products:
        values = []
        for v in product["row"].values():
            # Export user-provided text as text, never as executable spreadsheet formulas.
            values.append("'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v)
        sheet.append(values)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="167B64")
        cell.font = Font(color="FFFFFF", bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(50, max(16, len(str(column[0].value)) * 2))
    notes = book.create_sheet("안내")
    notes.append(["추정월매출은 월구매자수 × 판매가 × 2이며 제품 간 규모 비교용입니다."])
    notes.append(["단가기준이 다른 행의 단가를 직접 비교하지 마세요."])
    notes.append(["업로드 데이터는 브라우저 탭에서만 사용되며 원본 저장소를 변경하지 않습니다."])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def simulate(payload):
    if not isinstance(payload, dict):
        raise ValueError("JSON 객체가 필요합니다.")
    index = payload.get("index")
    price = payload.get("price")
    if type(index) is not int or not 0 <= index < len(load_specs()):
        raise ValueError("규격을 선택하세요.")
    if type(price) is not int or not 100 <= price <= 10000000:
        raise ValueError("목표 판매가는 100~10,000,000원 사이의 정수로 입력하세요.")
    spec = load_specs()[index]
    capacity, unit = parse_capacity(spec["spec"].replace(",", ""))
    scale, _ = unit_price_basis(unit)
    target_unit = price / (capacity * spec["qty"]) * scale
    return {"price": price, "cost": round(price / 3.5), "costEach": round(price / 3.5 / spec["qty"]),
            "unit": round(target_unit, 1), "basis": spec["basis"],
            "benchDiff": round((target_unit / spec["benchUnit"] - 1) * 100, 1)}


class Handler(BaseHTTPRequestHandler):
    server_version = "ResearchDashboard/1.0"

    def send(self, status, body, content_type="application/json; charset=utf-8", filename=None):
        if not isinstance(body, bytes):
            body = json.dumps(body, ensure_ascii=False, allow_nan=False).encode()
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "same-origin")
        self.send_header("Content-Security-Policy", "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline' https://fonts.googleapis.com; font-src 'self' https://fonts.gstatic.com; img-src 'self' data:; connect-src 'self'; base-uri 'none'; frame-ancestors 'none'")
        if filename:
            self.send_header("Content-Disposition", f"attachment; filename*=UTF-8''{quote(filename)}")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path = urlsplit(self.path).path
        if path == "/api/dashboard":
            return self.send(200, snapshot(load_products()))
        if path == "/api/export.xlsx":
            return self.send(200, make_workbook(load_products()), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "market_research.xlsx")
        if path == "/api/document":
            return self.send(200, (ROOT / "신제품_생산견적요청서.docx").read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "product_quote.docx")
        if path == "/api/specs.json":
            return self.send(200, json.dumps(load_specs(), ensure_ascii=False, indent=2).encode(), "application/json; charset=utf-8", "target_specs.json")
        if path == "/api/template.json":
            return self.send(200, (ROOT / "data/extracted/세탁세제.json").read_bytes(), "application/json; charset=utf-8", "extraction_example.json")
        if path == "/health":
            return self.send(200, {"status": "ok"})
        # Explicit allowlist: never expose the repository, .git, or arbitrary paths.
        assets = {"/": ("index.html", "text/html; charset=utf-8"), "/app.js": ("app.js", "text/javascript; charset=utf-8"),
                  "/styles.css": ("styles.css", "text/css; charset=utf-8"), "/favicon.svg": ("favicon.svg", "image/svg+xml")}
        if path in assets:
            name, content_type = assets[path]
            file = Path(__file__).parent / name
            if file.exists():
                return self.send(200, file.read_bytes(), content_type)
        self.send(404, {"error": "요청한 파일을 찾을 수 없습니다."})

    def do_POST(self):
        if urlsplit(self.path).path not in ("/api/import", "/api/export.xlsx", "/api/simulate"):
            return self.send(404, {"error": "지원하지 않는 요청입니다."})
        if self.headers.get_content_type() != "application/json":
            return self.send(415, {"error": "application/json 형식으로 전송하세요."})
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if not 0 < length <= MAX_BODY:
                return self.send(413, {"error": "최대 4MB까지 전송할 수 있습니다."})
            payload = json.loads(self.rfile.read(length))
            path = urlsplit(self.path).path
            if path == "/api/simulate":
                return self.send(200, simulate(payload))
            products = products_from_payload(payload)
            if path == "/api/import":
                return self.send(200, snapshot(products, True))
            return self.send(200, make_workbook(products), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "market_research.xlsx")
        except (ValueError, TypeError, OverflowError) as exc:
            self.send(400, {"error": str(exc)})
        except Exception as exc:
            self.log_error("Unexpected processing error: %s", type(exc).__name__)
            self.send(500, {"error": "처리 중 오류가 발생했습니다. 데이터 형식을 확인하세요."})


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--port", type=int, default=3000)
    parser.add_argument("--host", default="127.0.0.1")
    args = parser.parse_args()
    server = ThreadingHTTPServer((args.host, args.port), Handler)
    print(f"Dashboard: http://{args.host}:{args.port}", flush=True)
    server.serve_forever()


if __name__ == "__main__":
    main()
