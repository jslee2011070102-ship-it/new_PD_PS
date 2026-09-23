import importlib.util
from io import BytesIO
from pathlib import Path
import unittest
from openpyxl import load_workbook

spec = importlib.util.spec_from_file_location("server", Path(__file__).resolve().parents[1] / "dashboard/server.py")
server = importlib.util.module_from_spec(spec)
spec.loader.exec_module(server)


class DashboardTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.products = server.load_products()

    def test_real_dataset(self):
        summary = server.summarize(self.products)
        self.assertEqual(summary["count"], 150)
        self.assertEqual(summary["validation"], {"일치": 145, "표기없음": 5})
        self.assertEqual(len(server.load_specs()), 12)
        self.assertEqual([c["count"] for c in summary["categories"]], [25] * 6)

    def test_units_remain_separate(self):
        bases = {p["basis"] for p in self.products}
        self.assertEqual(bases, {"100ml당", "100g당", "1개당"})
        self.assertTrue(all(p["basis"] == "1개당" for p in self.products if p["category"] == "캡슐세제"))

    def test_satisfaction_is_not_buyers(self):
        info = server.ProductInfo.model_validate({k: v for k, v in self.products[0].items() if k in server.ProductInfo.model_fields})
        info.monthly_buyers_text = "한 달간 2만명 이상 만족했어요"
        row = server.row_from_info(info, "test.jpg", {"survey_date": "2026-09-21", "channel": "쿠팡", "category": "세탁세제"})
        self.assertIsNone(row["월구매자수(명)"])
        self.assertIsNone(row["추정월매출(원)"])

    def test_import_validation(self):
        for entries in ([], {}, [None], [{"brand": "incomplete"}]):
            with self.assertRaises(ValueError):
                server.normalize("세탁세제", entries)
        with self.assertRaises(ValueError):
            server.normalize("unknown", [{}])
        with self.assertRaises(ValueError):
            server.products_from_payload({"overrides": {"../bad": []}})

    def test_export_is_real_excel_and_escapes_formulas(self):
        import copy
        products = copy.deepcopy(self.products)
        products[0]["row"]["제품명"] = "=1+1"
        book = load_workbook(BytesIO(server.make_workbook(products)))
        self.assertEqual(book["전체"].max_row, 151)
        self.assertEqual(book["전체"]["F2"].data_type, "s")
        self.assertEqual(book["전체"].freeze_panes, "A2")

    def test_cost_simulator(self):
        result = server.simulate({"index": 0, "price": 14000})
        self.assertEqual(result["cost"], 4000)
        self.assertEqual(result["costEach"], 1000)
        for price in (-1, 0, "14000", True, 10000001):
            with self.assertRaises(ValueError):
                server.simulate({"index": 0, "price": price})


if __name__ == "__main__":
    unittest.main()
